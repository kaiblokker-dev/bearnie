"""Opent een Excel-export van de meettool met openpyxl en geeft kerngegevens terug als JSON."""
import json
import sys

import openpyxl

wb = openpyxl.load_workbook(sys.argv[1])
uit = {"bladen": wb.sheetnames}


def datarijen(ws):
    rijen = list(ws.iter_rows(values_only=True))
    kop = next(i for i, r in enumerate(rijen) if r and r[0] in ("MetingID", "FrequentieID"))
    return rijen[kop], [r for r in rijen[kop + 1:] if r and r[0]]


for naam, sleutel in (("Procesmetingen", "procesmetingen"), ("Stapmetingen", "stapmetingen"), ("Frequentie", "frequentie")):
    kop, rijen = datarijen(wb[naam])
    uit[sleutel] = len(rijen)

# Ruwe tabbladen mogen geen berekende kolommen bevatten
ruwe_koppen = [c for naam in ("Procesmetingen", "Stapmetingen", "Frequentie") for c in datarijen(wb[naam])[0]]
uit["ruw_zonder_berekening"] = not any(k and ("Mediaan" in k or "Totale actieve" in k or "Totale wacht" in k or " per " in k) for k in ruwe_koppen)

res = list(wb["Resultaten"].iter_rows(values_only=True))
kop_index = next(i for i, r in enumerate(res) if r and r[0] == "ProcesID")
kop = res[kop_index]
for r in res[kop_index + 1:]:
    if r and r[0] == "PR24" and r[2] == "Normaal":
        uit["pr24_normaal_mediaan"] = r[kop.index("Mediaan actieve tijd per uitvoering (min) – primair")]
        uit["pr24_normaal_mediaan_omvang"] = r[kop.index("Mediaan actieve tijd per omvangseenheid (min) – aanvullend")]
        uit["pr24_belasting"] = r[kop.index("Geschatte actieve tijd per meetperiode (min)")]
        uit["pr24_eenheden"] = [r[kop.index("Uitvoeringseenheid")], r[kop.index("Omvangseenheid")]]
        uit["pr24_ids"] = r[kop.index("Onderliggende MetingID's")]
        break

pm_kop = datarijen(wb["Procesmetingen"])[0]
fr_kop = datarijen(wb["Frequentie"])[0]
uit["ruw_eenheden"] = all(k in pm_kop for k in ("Aantal uitvoeringen", "Uitvoeringseenheid", "Omvang", "Omvangseenheid", "Herkomst aantal uitvoeringen")) and all(
    k in fr_kop for k in ("Aantal uitvoeringen", "Uitvoeringseenheid", "Totale omvang", "Omvangseenheid"))

methode = " ".join(str(c) for r in wb["Methode"].iter_rows(values_only=True) for c in r if c)
uit["methode_eenheden"] = all(t in methode for t in ("Uitvoeringseenheid", "Omvangseenheid", "PRIMAIRE UITKOMST", "Eenheden per proces", "dossier (meervoud: dossiers)", "dienstperiode (meervoud: dienstperioden)"))
uit["methode_ontbrekend"] = "NIET automatisch als nul verwerkt" in methode and "Versienummer tool" in methode

print(json.dumps(uit))
